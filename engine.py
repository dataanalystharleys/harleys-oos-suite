import os, shutil, json, io, datetime, string
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from PIL import Image, ImageDraw, ImageFont

# Base Directories
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
if os.path.exists(r"C:\Users\harle\Downloads\oos"):
    DOWNLOADS_DIR = r"C:\Users\harle\Downloads"
    OOS_ROOT = os.path.join(DOWNLOADS_DIR, "oos")
else:
    OOS_ROOT = BASE_DIR
    DOWNLOADS_DIR = os.path.join(BASE_DIR, "downloads")
    os.makedirs(DOWNLOADS_DIR, exist_ok=True)

TIMINGS_DB_PATH = os.path.join(OOS_ROOT, "outlet_timings_db.json")

# Master Reference Paths
MASTER_TEMPLATE_PATH = os.path.join(OOS_ROOT, "Oct-2026 OOS SUMMARY.xlsx")
if not os.path.exists(MASTER_TEMPLATE_PATH):
    MASTER_TEMPLATE_PATH = os.path.join(OOS_ROOT, "Oct - 03", "Oct-2026 OOS SUMMARY.xlsx")
if not os.path.exists(MASTER_TEMPLATE_PATH):
    MASTER_TEMPLATE_PATH = os.path.join(DOWNLOADS_DIR, "Oct-2026 OOS SUMMARY (1).xlsx")

# Excel Styles
font_banner = Font(name="Calibri", size=11, bold=True, color="000000")
font_hdr = Font(name="Calibri", size=11, bold=True, color="000000")
font_cell = Font(name="Calibri", size=10, bold=False, color="1A1A1A")
font_cell_bold = Font(name="Calibri", size=10, bold=True, color="1A1A1A")
font_out = Font(name="Calibri", size=10, bold=True, color="C00000")
font_in = Font(name="Calibri", size=10, bold=True, color="276A3C")

fill_banner = PatternFill(start_color="E6F2FF", end_color="E6F2FF", fill_type="solid")
fill_hdr = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
fill_item = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
fill_out = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
fill_in = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")

thin_side = Side(border_style="thin", color="D3D3D3")
cell_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

align_center = Alignment(horizontal="center", vertical="center")
align_left = Alignment(horizontal="left", vertical="center")

CITY_ORDER = {'AP': 1, 'BLR': 2, 'HYD': 3, 'MUM': 4, 'NCR': 5, 'PUN': 6}

def get_city(branch_name):
    b = str(branch_name).upper()
    if 'VIJ' in b or b.startswith('AP-'): return 'AP'
    if 'BLR' in b or b.startswith('KA-'): return 'BLR'
    if 'HYD' in b or b.startswith('TG-'): return 'HYD'
    if 'MUM' in b or b.startswith('MH-MUM'): return 'MUM'
    if 'PUN' in b or b.startswith('MH-PUN'): return 'PUN'
    if 'NCR' in b or 'GUR' in b or 'DELHI' in b: return 'NCR'
    return 'OTHER'

def load_timings_db():
    if os.path.exists(TIMINGS_DB_PATH):
        with open(TIMINGS_DB_PATH, 'r') as f:
            return json.load(f)
    return {}

def save_timings_db(db):
    os.makedirs(OOS_ROOT, exist_ok=True)
    with open(TIMINGS_DB_PATH, 'w') as f:
        json.dump(db, f, indent=2)

def get_outlet_opening_time(branch, date_obj, timings_db=None):
    if timings_db is None:
        timings_db = load_timings_db()
    day_name = date_obj.strftime('%A')
    info = timings_db.get(branch)
    if info:
        return info.get(day_name) or info.get('default') or '09:00'
    return '09:00'

def load_mappings_and_history():
    wb_ref = openpyxl.load_workbook(MASTER_TEMPLATE_PATH, read_only=True)
    supp_lower = {str(r[0]).strip().lower(): str(r[1]).strip() for idx, r in enumerate(wb_ref['Supporting'].iter_rows(values_only=True)) if idx > 0 and r[0] and r[1]}
    top20 = {(str(r[0]).strip(), str(r[2]).strip()): int(r[4]) for idx, r in enumerate(wb_ref['Top 20'].iter_rows(values_only=True)) if idx > 0 and r[0] and r[2] and r[5] == 1}
    
    ws_raw = wb_ref['SOLD OUT RAW']
    past_events = []
    for idx, r in enumerate(ws_raw.iter_rows(values_only=True)):
        if idx == 0: continue
        raw_b = str(r[11]).strip() if r[11] else ''
        odoo = supp_lower.get(raw_b.lower(), raw_b)
        item = str(r[1]).strip() if r[1] else ''
        item_io = str(r[6]).strip().upper() if r[6] else ''
        dt = r[9]
        if (odoo, item) in top20 and isinstance(dt, datetime.datetime):
            past_events.append({'odoo': odoo, 'item': item, 'in_out': item_io, 'dt': dt})
    wb_ref.close()
    return supp_lower, top20, past_events

# -------------------------------------------------------------
# SCREENSHOT RENDERING HELPER (PIL)
# -------------------------------------------------------------
SCALE = 2
FONT_REGULAR = "C:/Windows/Fonts/calibri.ttf"
FONT_BOLD = "C:/Windows/Fonts/calibrib.ttf"

def s(val):
    return int(val * SCALE)

def get_font(path, size):
    scaled_size = int(size * SCALE)
    if os.path.exists(path):
        try:
            return ImageFont.truetype(path, scaled_size)
        except Exception:
            pass
    for fb in ["DejaVuSans-Bold.ttf" if "bold" in path.lower() else "DejaVuSans.ttf", "arial.ttf", "Calibri.ttf"]:
        try:
            return ImageFont.truetype(fb, scaled_size)
        except Exception:
            pass
    return ImageFont.load_default()

f_header = get_font(FONT_BOLD, 11)
f_row_hdr = get_font(FONT_REGULAR, 9)
f_body = get_font(FONT_REGULAR, 10.5)
f_body_bold = get_font(FONT_BOLD, 10.5)
f_kpi = get_font(FONT_BOLD, 11)
f_status_bold = get_font(FONT_BOLD, 10)

C_EXCEL_HEADER_BG = "#E8E8E8"
C_EXCEL_HEADER_BORDER = "#D4D4D4"
C_EXCEL_HEADER_TEXT = "#555555"
C_NAVY = "#1B365D"
C_NAVY_BORDER = "#142846"
C_GRID = "#D9D9D9"
C_DATE_BG = "#FFF2CC"
C_DATE_BORDER = "#D6B656"
C_WHITE = "#FFFFFF"
C_BLACK = "#1A1A1A"
C_OUT_BG = "#FCE4D6"
C_OUT_TXT = "#C00000"
C_IN_BG = "#E2EFDA"
C_IN_TXT = "#276A3C"

def render_table_image(date_str, title_banner, cols, display_rows, out_paths, is_evening=False):
    row_hdr_w = 40
    col_hdr_h = 24
    kpi_row_h = 28
    table_hdr_h = 32
    data_row_h = 24

    total_w = row_hdr_w + sum(c[1] for c in cols)
    total_h = col_hdr_h + kpi_row_h + table_hdr_h + len(display_rows) * data_row_h + 8

    img = Image.new("RGB", (s(total_w), s(total_h)), "#FFFFFF")
    draw = ImageDraw.Draw(img)

    draw.rectangle([0, 0, s(row_hdr_w), s(col_hdr_h)], fill=C_EXCEL_HEADER_BG, outline=C_EXCEL_HEADER_BORDER, width=s(1))

    cur_x = row_hdr_w
    col_letters = list(string.ascii_uppercase)[:len(cols)]
    for idx, (name, width, _) in enumerate(cols):
        x1, y1 = s(cur_x), 0
        x2, y2 = s(cur_x + width), s(col_hdr_h)
        draw.rectangle([x1, y1, x2, y2], fill=C_EXCEL_HEADER_BG, outline=C_EXCEL_HEADER_BORDER, width=s(1))
        letter = col_letters[idx]
        bbox = f_row_hdr.getbbox(letter)
        tx = x1 + (x2 - x1 - (bbox[2] - bbox[0])) // 2
        ty = y1 + (y2 - y1 - (bbox[3] - bbox[1])) // 2
        draw.text((tx, ty), letter, fill=C_EXCEL_HEADER_TEXT, font=f_row_hdr)
        cur_x += width

    cur_y = col_hdr_h
    draw.rectangle([0, s(cur_y), s(row_hdr_w), s(cur_y + kpi_row_h)], fill=C_EXCEL_HEADER_BG, outline=C_EXCEL_HEADER_BORDER, width=s(1))
    bbox = f_row_hdr.getbbox("1")
    draw.text(((s(row_hdr_w) - (bbox[2] - bbox[0])) // 2, s(cur_y) + (s(kpi_row_h) - (bbox[3] - bbox[1])) // 2), "1", fill=C_EXCEL_HEADER_TEXT, font=f_row_hdr)

    c_a_w = cols[0][1] + cols[1][1]
    bx1 = s(row_hdr_w)
    bx2 = bx1 + s(c_a_w)
    draw.rectangle([bx1, s(cur_y), bx2, s(cur_y + kpi_row_h)], fill=C_DATE_BG, outline=C_DATE_BORDER, width=s(1))
    bbox = f_kpi.getbbox(date_str)
    draw.text((bx1 + (bx2 - bx1 - (bbox[2] - bbox[0])) // 2, s(cur_y) + (s(kpi_row_h) - (bbox[3] - bbox[1])) // 2), date_str, fill=C_BLACK, font=f_kpi)

    kpi_x1 = bx2
    kpi_x2 = s(row_hdr_w + sum(c[1] for c in cols))
    draw.rectangle([kpi_x1, s(cur_y), kpi_x2, s(cur_y + kpi_row_h)], fill=C_WHITE, outline=C_GRID, width=s(1))
    draw.text((kpi_x1 + s(14), s(cur_y) + s(6)), title_banner, fill=C_BLACK, font=f_kpi)

    cur_y += kpi_row_h
    draw.rectangle([0, s(cur_y), s(row_hdr_w), s(cur_y + table_hdr_h)], fill=C_EXCEL_HEADER_BG, outline=C_EXCEL_HEADER_BORDER, width=s(1))
    bbox = f_row_hdr.getbbox("2")
    draw.text(((s(row_hdr_w) - (bbox[2] - bbox[0])) // 2, s(cur_y) + (s(table_hdr_h) - (bbox[3] - bbox[1])) // 2), "2", fill=C_EXCEL_HEADER_TEXT, font=f_row_hdr)

    cur_x = row_hdr_w
    for name, width, align in cols:
        x1, y1 = s(cur_x), s(cur_y)
        x2, y2 = s(cur_x + width), s(cur_y + table_hdr_h)
        draw.rectangle([x1, y1, x2, y2], fill=C_NAVY, outline=C_NAVY_BORDER, width=s(1))
        
        words = name.split()
        if len(words) > 3 and width < 200:
            line1 = " ".join(words[:2]); line2 = " ".join(words[2:])
            bb1 = f_header.getbbox(line1); bb2 = f_header.getbbox(line2)
            h1 = bb1[3] - bb1[1]; h2 = bb2[3] - bb2[1]
            tot_h = h1 + h2 + s(3)
            ty = y1 + (y2 - y1 - tot_h) // 2
            tx1 = x1 + (x2 - x1 - (bb1[2] - bb1[0])) // 2
            tx2 = x1 + (x2 - x1 - (bb2[2] - bb2[0])) // 2
            draw.text((tx1, ty), line1, fill=C_WHITE, font=f_header)
            draw.text((tx2, ty + h1 + s(3)), line2, fill=C_WHITE, font=f_header)
        else:
            bbox = f_header.getbbox(name)
            tx = x1 + (x2 - x1 - (bbox[2] - bbox[0])) // 2
            ty = y1 + (y2 - y1 - (bbox[3] - bbox[1])) // 2
            draw.text((tx, ty), name, fill=C_WHITE, font=f_header)
        cur_x += width

    for r_idx, row_vals in enumerate(display_rows):
        row_num_str = str(r_idx + 3)
        cur_y += data_row_h
        
        draw.rectangle([0, s(cur_y), s(row_hdr_w), s(cur_y + data_row_h)], fill=C_EXCEL_HEADER_BG, outline=C_EXCEL_HEADER_BORDER, width=s(1))
        bbox = f_row_hdr.getbbox(row_num_str)
        draw.text(((s(row_hdr_w) - (bbox[2] - bbox[0])) // 2, s(cur_y) + (s(data_row_h) - (bbox[3] - bbox[1])) // 2), row_num_str, fill=C_EXCEL_HEADER_TEXT, font=f_row_hdr)
        
        cur_x = row_hdr_w
        for c_idx, val in enumerate(row_vals):
            width = cols[c_idx][1]; align = cols[c_idx][2]
            x1, y1 = s(cur_x), s(cur_y); x2, y2 = s(cur_x + width), s(cur_y + data_row_h)
            
            if is_evening and c_idx == 6:
                bg = C_OUT_BG if val == 'OUT' else C_IN_BG
                draw.rectangle([x1, y1, x2, y2], fill=bg, outline=C_GRID, width=s(1))
            else:
                draw.rectangle([x1, y1, x2, y2], fill=C_WHITE, outline=C_GRID, width=s(1))
                
            if val != "":
                if is_evening and c_idx == 6:
                    fnt = f_status_bold
                    txt_color = C_OUT_TXT if val == 'OUT' else C_IN_TXT
                else:
                    fnt = f_body_bold if c_idx == 0 or c_idx == 2 or (c_idx == 3 and val != "") else f_body
                    txt_color = C_BLACK
                    
                bbox = fnt.getbbox(str(val))
                tw = bbox[2] - bbox[0]; th = bbox[3] - bbox[1]
                if align == "center": tx = x1 + (x2 - x1 - tw) // 2
                elif align == "right": tx = x2 - tw - s(10)
                else: tx = x1 + s(8)
                ty = y1 + (y2 - y1 - th) // 2
                draw.text((tx, ty), str(val), fill=txt_color, font=fnt)
            cur_x += width

    for p in out_paths:
        os.makedirs(os.path.dirname(p), exist_ok=True)
        img.save(p, "PNG")

# -------------------------------------------------------------
# CORE PROCESSING PIPELINE
# -------------------------------------------------------------
def run_oos_pipeline(
    csv_input,
    target_date=None,
    run_morning=True,
    run_evening=True,
    progress_cb=None
):
    def log(msg):
        if progress_cb: progress_cb(msg)
        print(msg)

    log("Loading mappings, Top 20 flags, and timing rules...")
    supp_lower, top20, past_events = load_mappings_and_history()
    timings_db = load_timings_db()

    if isinstance(csv_input, str):
        df_raw = pd.read_csv(csv_input)
    else:
        df_raw = pd.read_csv(csv_input)

    df_raw['dt'] = pd.to_datetime(df_raw['Date and Time'])
    
    if target_date is None:
        target_date = df_raw['dt'].dt.date.mode()[0]
    elif isinstance(target_date, str):
        target_date = pd.to_datetime(target_date).date()

    day_name = target_date.strftime('%A')
    date_str_slash = target_date.strftime('%Y/%m/%d')
    date_str_dash = target_date.strftime('%Y-%m-%d')
    date_folder_name = f"Oct - {target_date.strftime('%d')}" if target_date.month == 10 else f"{target_date.strftime('%b - %d')}"
    
    log(f"Processing OOS for Date: {date_str_dash} ({day_name})...")

    date_dir = os.path.join(OOS_ROOT, date_folder_name)
    morning_dir = os.path.join(date_dir, "Morning")
    evening_dir = os.path.join(date_dir, "Evening")
    os.makedirs(date_dir, exist_ok=True)

    raw_save_path = os.path.join(date_dir, f"{target_date.strftime('%d-%b-%Y')} - Item sold out change history.csv")
    if isinstance(csv_input, str) and os.path.exists(csv_input):
        if os.path.abspath(csv_input) != os.path.abspath(raw_save_path):
            shutil.copy2(csv_input, raw_save_path)
    else:
        df_raw.to_csv(raw_save_path, index=False)

    combined_events = list(past_events)
    for idx, r in df_raw.iterrows():
        raw_b = str(r['Branch Name']).strip()
        odoo = supp_lower.get(raw_b.lower(), raw_b)
        item = str(r['Item Name']).strip()
        item_io = str(r['Item Item In/Out' if 'Item Item In/Out' in df_raw.columns else 'Item In/Out']).strip().upper()
        dt = r['dt']
        if (odoo, item) in top20:
            combined_events.append({'odoo': odoo, 'item': item, 'in_out': item_io, 'dt': dt.to_pydatetime()})

    combined_events.sort(key=lambda x: x['dt'])
    
    by_pair = {}
    for e in combined_events:
        by_pair.setdefault((e['odoo'], e['item']), []).append(e)

    results = {
        'date_str': date_str_slash,
        'date_dash': date_str_dash,
        'day_name': day_name,
        'date_dir': date_dir,
        'morning': None,
        'evening': None
    }

    # 1. MORNING
    morning_list = []
    if run_morning:
        log("Evaluating Morning OOS status at outlet opening hours...")
        for (odoo, item), evs in by_pair.items():
            city = get_city(odoo)
            rank = top20[(odoo, item)]
            op_time_str = get_outlet_opening_time(odoo, target_date, timings_db)
            oh, om = map(int, op_time_str.split(':'))
            open_dt = datetime.datetime(target_date.year, target_date.month, target_date.day, oh, om, 0)
            
            evs_before_open = [x for x in evs if x['dt'] <= open_dt]
            evs_target = [x for x in evs if x['dt'].date() == target_date]
            
            is_morning_oos = False
            resolved_in_str = None
            
            if evs_before_open:
                if evs_before_open[-1]['in_out'] == 'OUT':
                    is_morning_oos = True
                    after_open = [x for x in evs if x['dt'] > open_dt and x['in_out'] == 'IN']
                    if after_open:
                        resolved_in_str = after_open[0]['dt'].strftime('%H:%M')
                    else:
                        resolved_in_str = 'Unresolved'
            elif evs_target:
                if evs_target[0]['in_out'] == 'IN' and evs_target[0]['dt'] > open_dt:
                    is_morning_oos = True
                    resolved_in_str = evs_target[0]['dt'].strftime('%H:%M')
                    
            if is_morning_oos:
                morning_list.append({
                    'date': date_str_slash,
                    'city': city,
                    'odoo': odoo,
                    'open_time': op_time_str,
                    'rank': rank,
                    'item': item,
                    'resolved_in': resolved_in_str
                })

        morning_list.sort(key=lambda x: (CITY_ORDER.get(x['city'], 99), x['odoo'], x['rank']))
        os.makedirs(morning_dir, exist_ok=True)
        
        m_json_path = os.path.join(morning_dir, f"morning_oos_data_{date_str_dash}.json")
        with open(m_json_path, 'w') as f: json.dump(morning_list, f, indent=2)

        log("Generating Morning multi-tab Excel workbook...")
        wb_m = openpyxl.Workbook()
        ws_all_m = wb_m.active; ws_all_m.title = "ALL CITIES"
        populate_morning_excel_sheet(ws_all_m, date_str_slash, morning_list)
        for c in ['AP', 'BLR', 'HYD', 'MUM', 'PUN']:
            c_items = [x for x in morning_list if x['city'] == c]
            if c_items:
                ws_c = wb_m.create_sheet(title=c)
                populate_morning_excel_sheet(ws_c, date_str_slash, c_items)
        m_xlsx_path = os.path.join(morning_dir, f"Morning_OOS_Summary_{date_str_dash}.xlsx")
        wb_m.save(m_xlsx_path); wb_m.close()

        log("Rendering high-resolution Morning screenshot images...")
        m_cols = [
            ("S/No", 50, "center"), ("Date", 85, "center"), ("CITY", 60, "center"),
            ("Branch Name", 280, "left"), ("Outlet Timing", 100, "center"),
            ("Top 20 Rank", 95, "center"), ("Item Name", 380, "left")
        ]
        m_all_rows = build_morning_display_rows(morning_list)
        m_all_banner = f"MORNING OOS REPORT (OPENING)   |   Total Outlets: {len(set(x['odoo'] for x in morning_list))}   |   Total Items with OOS: {len(morning_list)}"
        m_png_all = os.path.join(morning_dir, f"Morning_OOS_Summary_ALL_CITIES_{date_str_dash}.png")
        m_png_all_city = os.path.join(date_dir, "ALL CITIES", f"Morning_OOS_Summary_ALL_CITIES_{date_str_dash}.png")
        render_table_image(date_str_slash, m_all_banner, m_cols, m_all_rows, [m_png_all, m_png_all_city], is_evening=False)

        m_city_pngs = {}
        for c in ['AP', 'BLR', 'HYD', 'MUM', 'PUN']:
            c_items = [x for x in morning_list if x['city'] == c]
            if c_items:
                c_rows = build_morning_display_rows(c_items)
                c_banner = f"MORNING OOS [{c}]   |   Total Outlets: {len(set(x['odoo'] for x in c_items))}   |   Total Items with OOS: {len(c_items)}"
                c_png = os.path.join(morning_dir, f"Morning_OOS_Summary_{c}_{date_str_dash}.png")
                c_png_city = os.path.join(date_dir, c, f"Morning_OOS_Summary_{c}_{date_str_dash}.png")
                render_table_image(date_str_slash, c_banner, m_cols, c_rows, [c_png, c_png_city], is_evening=False)
                m_city_pngs[c] = c_png

        results['morning'] = {
            'total_items': len(morning_list),
            'total_outlets': len(set(x['odoo'] for x in morning_list)),
            'items': morning_list,
            'excel_path': m_xlsx_path,
            'json_path': m_json_path,
            'png_all': m_png_all,
            'city_pngs': m_city_pngs
        }

    # 2. EVENING
    evening_combined = []
    if run_evening:
        log("Evaluating Evening OOS status (5:00 PM cutoff with Status field)...")
        cutoff_dt = datetime.datetime(target_date.year, target_date.month, target_date.day, 17, 0, 0)
        
        e_dict_items = {}
        for (odoo, item), evs in by_pair.items():
            evs_up_to_cutoff = [x for x in evs if x['dt'] <= cutoff_dt]
            if evs_up_to_cutoff and evs_up_to_cutoff[-1]['in_out'] == 'OUT':
                e_dict_items[(odoo, item)] = evs_up_to_cutoff[-1]

        m_dict_items = {(x['odoo'], x['item']): x for x in morning_list}
        all_e_keys = set(list(m_dict_items.keys()) + list(e_dict_items.keys()))

        for (odoo, item) in all_e_keys:
            m_info = m_dict_items.get((odoo, item))
            e_ev = e_dict_items.get((odoo, item))
            city = m_info['city'] if m_info else get_city(odoo)
            rank = m_info['rank'] if m_info else top20[(odoo, item)]
            
            if e_ev is not None:
                status = 'OUT'
                note = 'Still OUT' if m_info else f"Went OUT at {e_ev['dt'].strftime('%H:%M')}"
            else:
                status = 'IN'
                note = f"Restored IN at {m_info.get('resolved_in', 'Day') if m_info else 'Day'}"
                
            evening_combined.append({
                'date': date_str_slash,
                'city': city,
                'odoo': odoo,
                'rank': rank,
                'item': item,
                'status': status,
                'note': note
            })

        evening_combined.sort(key=lambda x: (CITY_ORDER.get(x['city'], 99), x['odoo'], x['rank']))
        os.makedirs(evening_dir, exist_ok=True)

        e_json_path = os.path.join(evening_dir, f"evening_oos_data_{date_str_dash}.json")
        with open(e_json_path, 'w') as f: json.dump(evening_combined, f, indent=2)

        log("Generating Evening multi-tab Excel workbook...")
        wb_e = openpyxl.Workbook()
        ws_all_e = wb_e.active; ws_all_e.title = "ALL CITIES"
        populate_evening_excel_sheet(ws_all_e, date_str_slash, evening_combined)
        for c in ['AP', 'BLR', 'HYD', 'MUM', 'PUN']:
            c_items = [x for x in evening_combined if x['city'] == c]
            if c_items:
                ws_c = wb_e.create_sheet(title=c)
                populate_evening_excel_sheet(ws_c, date_str_slash, c_items)
        e_xlsx_path = os.path.join(evening_dir, f"Evening_OOS_Summary_{date_str_dash}.xlsx")
        wb_e.save(e_xlsx_path); wb_e.close()

        log("Rendering high-resolution Evening screenshot images...")
        e_cols = [
            ("S/No", 50, "center"), ("Date", 85, "center"), ("CITY", 60, "center"),
            ("Branch Name", 280, "left"), ("Top 20 Rank", 95, "center"),
            ("Item Name", 380, "left"), ("Status", 75, "center")
        ]
        e_all_rows = build_evening_display_rows(evening_combined)
        out_tot = sum(1 for x in evening_combined if x['status'] == 'OUT')
        in_tot = sum(1 for x in evening_combined if x['status'] == 'IN')
        e_all_banner = f"EVENING OOS STATUS (5:00 PM CUTOFF)   |   Total: {len(evening_combined)}   |   Still OUT: {out_tot}   |   Restored IN: {in_tot}"
        e_png_all = os.path.join(evening_dir, f"Evening_OOS_Summary_ALL_CITIES_{date_str_dash}.png")
        e_png_all_city = os.path.join(date_dir, "ALL CITIES", f"Evening_OOS_Summary_ALL_CITIES_{date_str_dash}.png")
        render_table_image(date_str_slash, e_all_banner, e_cols, e_all_rows, [e_png_all, e_png_all_city], is_evening=True)

        e_city_pngs = {}
        for c in ['AP', 'BLR', 'HYD', 'MUM', 'PUN']:
            c_items = [x for x in evening_combined if x['city'] == c]
            if c_items:
                c_rows = build_evening_display_rows(c_items)
                c_out = sum(1 for x in c_items if x['status'] == 'OUT')
                c_in = sum(1 for x in c_items if x['status'] == 'IN')
                c_banner = f"EVENING OOS STATUS [{c}] (5:00 PM CUTOFF)   |   Total: {len(c_items)}   |   Still OUT: {c_out}   |   Restored IN: {c_in}"
                c_png = os.path.join(evening_dir, f"Evening_OOS_Summary_{c}_{date_str_dash}.png")
                c_png_city = os.path.join(date_dir, c, f"Evening_OOS_Summary_{c}_{date_str_dash}.png")
                render_table_image(date_str_slash, c_banner, e_cols, c_rows, [c_png, c_png_city], is_evening=True)
                e_city_pngs[c] = c_png

        results['evening'] = {
            'total_items': len(evening_combined),
            'total_outlets': len(set(x['odoo'] for x in evening_combined)),
            'still_out': out_tot,
            'restored_in': in_tot,
            'items': evening_combined,
            'excel_path': e_xlsx_path,
            'json_path': e_json_path,
            'png_all': e_png_all,
            'city_pngs': e_city_pngs
        }

    # 3. MASTER WORKBOOKS
    log("Updating master summary workbooks...")
    update_master_workbooks(date_dir, target_date, morning_list, evening_combined)

    log("OOS Pipeline execution completed successfully!")
    return results

def populate_morning_excel_sheet(ws, date_str, items):
    ws.cell(1, 1).value = ""
    ws.cell(1, 2).value = date_str
    ws.cell(1, 3).value = f"Total Outlets: {len(set(x['odoo'] for x in items))}"
    ws.cell(1, 4).value = f"Total Items with OOS: {len(items)}"
    for col in range(2, 8):
        c = ws.cell(1, col); c.fill = fill_banner; c.font = font_banner
        c.alignment = align_center if col == 2 else align_left
    ws.row_dimensions[1].height = 24
    
    headers = ['S/No', 'Date', 'CITY', 'Branch Name', 'Outlet Timing', 'Top 20 Rank', 'Item Name']
    for col_idx, h in enumerate(headers, 1):
        c = ws.cell(2, col_idx); c.value = h; c.font = font_hdr; c.fill = fill_hdr; c.alignment = align_center; c.border = cell_border
    ws.row_dimensions[2].height = 28
    
    curr_row = 3; s_no = 0; prev_b, prev_c = None, None
    for item in items:
        if item['odoo'] != prev_b:
            s_no += 1; b_disp, sno_disp, prev_b = item['odoo'], s_no, item['odoo']
        else: b_disp, sno_disp = "", ""
        if item['city'] != prev_c: c_disp, prev_c = item['city'], item['city']
        else: c_disp = ""
        
        ws.cell(curr_row, 1).value = sno_disp
        ws.cell(curr_row, 2).value = item['date']
        ws.cell(curr_row, 3).value = c_disp
        ws.cell(curr_row, 4).value = b_disp
        ws.cell(curr_row, 5).value = item['open_time']
        ws.cell(curr_row, 6).value = item['rank']
        ws.cell(curr_row, 7).value = item['item']
        
        ws.cell(curr_row, 1).alignment = align_center; ws.cell(curr_row, 1).font = font_cell_bold
        ws.cell(curr_row, 2).alignment = align_center; ws.cell(curr_row, 2).font = font_cell
        ws.cell(curr_row, 3).alignment = align_center; ws.cell(curr_row, 3).font = font_cell_bold
        ws.cell(curr_row, 4).alignment = align_left; ws.cell(curr_row, 4).font = font_cell_bold if b_disp else font_cell
        ws.cell(curr_row, 5).alignment = align_center; ws.cell(curr_row, 5).font = font_cell
        ws.cell(curr_row, 6).alignment = align_center; ws.cell(curr_row, 6).font = font_cell
        ws.cell(curr_row, 7).alignment = align_left; ws.cell(curr_row, 7).font = font_cell; ws.cell(curr_row, 7).fill = fill_item
        for ci in range(1, 8): ws.cell(curr_row, ci).border = cell_border
        ws.row_dimensions[curr_row].height = 20
        curr_row += 1
        
    for col_letter, width in {'A': 8, 'B': 14, 'C': 10, 'D': 42, 'E': 15, 'F': 14, 'G': 52}.items():
        ws.column_dimensions[col_letter].width = width

def populate_evening_excel_sheet(ws, date_str, items):
    out_cnt = sum(1 for x in items if x['status'] == 'OUT')
    in_cnt = sum(1 for x in items if x['status'] == 'IN')
    
    ws.cell(1, 1).value = ""
    ws.cell(1, 2).value = date_str
    ws.cell(1, 3).value = f"Total Outlets: {len(set(x['odoo'] for x in items))}"
    ws.cell(1, 4).value = f"Still OUT: {out_cnt}"
    ws.cell(1, 5).value = f"Restored IN: {in_cnt}"
    ws.cell(1, 6).value = "Cutoff: 5:00 PM (17:00)"
    for col in range(2, 8):
        c = ws.cell(1, col); c.fill = fill_banner; c.font = font_banner
        c.alignment = align_center if col == 2 else align_left
    ws.row_dimensions[1].height = 24
    
    headers = ['S/No', 'Date', 'CITY', 'Branch Name', 'Top 20 Rank', 'Item Name', 'Status']
    for col_idx, h in enumerate(headers, 1):
        c = ws.cell(2, col_idx); c.value = h; c.font = font_hdr; c.fill = fill_hdr; c.alignment = align_center; c.border = cell_border
    ws.row_dimensions[2].height = 28
    
    curr_row = 3; s_no = 0; prev_b, prev_c = None, None
    for item in items:
        if item['odoo'] != prev_b:
            s_no += 1; b_disp, sno_disp, prev_b = item['odoo'], s_no, item['odoo']
        else: b_disp, sno_disp = "", ""
        if item['city'] != prev_c: c_disp, prev_c = item['city'], item['city']
        else: c_disp = ""
        
        ws.cell(curr_row, 1).value = sno_disp
        ws.cell(curr_row, 2).value = item['date']
        ws.cell(curr_row, 3).value = c_disp
        ws.cell(curr_row, 4).value = b_disp
        ws.cell(curr_row, 5).value = item['rank']
        ws.cell(curr_row, 6).value = item['item']
        ws.cell(curr_row, 7).value = item['status']
        
        ws.cell(curr_row, 1).alignment = align_center; ws.cell(curr_row, 1).font = font_cell_bold
        ws.cell(curr_row, 2).alignment = align_center; ws.cell(curr_row, 2).font = font_cell
        ws.cell(curr_row, 3).alignment = align_center; ws.cell(curr_row, 3).font = font_cell_bold
        ws.cell(curr_row, 4).alignment = align_left; ws.cell(curr_row, 4).font = font_cell_bold if b_disp else font_cell
        ws.cell(curr_row, 5).alignment = align_center; ws.cell(curr_row, 5).font = font_cell
        ws.cell(curr_row, 6).alignment = align_left; ws.cell(curr_row, 6).font = font_cell; ws.cell(curr_row, 6).fill = fill_item
        
        ws.cell(curr_row, 7).alignment = align_center
        if item['status'] == 'OUT':
            ws.cell(curr_row, 7).font = font_out; ws.cell(curr_row, 7).fill = fill_out
        else:
            ws.cell(curr_row, 7).font = font_in; ws.cell(curr_row, 7).fill = fill_in
            
        for ci in range(1, 8): ws.cell(curr_row, ci).border = cell_border
        ws.row_dimensions[curr_row].height = 20
        curr_row += 1
        
    for col_letter, width in {'A': 8, 'B': 14, 'C': 10, 'D': 42, 'E': 14, 'F': 52, 'G': 14}.items():
        ws.column_dimensions[col_letter].width = width

def build_morning_display_rows(items):
    rows = []
    cur_b, cur_c, sno = None, None, 0
    for it in items:
        if it['odoo'] != cur_b:
            sno += 1; sno_str, b_str, cur_b = str(sno), it['odoo'], it['odoo']
        else: sno_str, b_str = "", ""
        if it['city'] != cur_c: c_str, cur_c = it['city'], it['city']
        else: c_str = ""
        rows.append((sno_str, it['date'], c_str, b_str, it['open_time'], str(it['rank']), it['item']))
    return rows

def build_evening_display_rows(items):
    rows = []
    cur_b, cur_c, sno = None, None, 0
    for it in items:
        if it['odoo'] != cur_b:
            sno += 1; sno_str, b_str, cur_b = str(sno), it['odoo'], it['odoo']
        else: sno_str, b_str = "", ""
        if it['city'] != cur_c: c_str, cur_c = it['city'], it['city']
        else: c_str = ""
        rows.append((sno_str, it['date'], c_str, b_str, str(it['rank']), it['item'], it['status']))
    return rows

def update_master_workbooks(date_dir, target_date, morning_items, evening_items):
    s_day = target_date.strftime('%d-%b')
    sname_m = f"MORNING OOS ({s_day})"
    sname_e = f"EVENING OOS ({s_day})"
    
    date_master = os.path.join(date_dir, "Oct-2026 OOS SUMMARY.xlsx")
    root_master = os.path.join(OOS_ROOT, "Oct-2026 OOS SUMMARY.xlsx")
    global_master = os.path.join(DOWNLOADS_DIR, "Oct-2026 OOS SUMMARY (1).xlsx")
    
    base_src = None
    for p in [date_master, root_master, global_master, MASTER_TEMPLATE_PATH]:
        if os.path.exists(p):
            base_src = p
            break
            
    if not base_src:
        return
        
    with open(base_src, 'rb') as f:
        f_bytes = io.BytesIO(f.read())
    wb = openpyxl.load_workbook(f_bytes)
    
    if morning_items:
        if sname_m in wb.sheetnames: del wb[sname_m]
        ws_m = wb.create_sheet(title=sname_m)
        populate_morning_excel_sheet(ws_m, target_date.strftime('%Y/%m/%d'), morning_items)
        
    if evening_items:
        if sname_e in wb.sheetnames: del wb[sname_e]
        ws_e = wb.create_sheet(title=sname_e)
        populate_evening_excel_sheet(ws_e, target_date.strftime('%Y/%m/%d'), evening_items)
        
    out_buf = io.BytesIO()
    wb.save(out_buf)
    wb.close()
    out_data = out_buf.getvalue()
    
    targets = [date_master, root_master]
    if os.path.exists(global_master):
        targets.append(global_master)
        
    for t_path in targets:
        try:
            with open(t_path, 'wb') as f:
                f.write(out_data)
        except Exception:
            pass

